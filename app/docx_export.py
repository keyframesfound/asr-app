"""Builds Word (.docx) downloads for the transcript and the AI summary.

The summary arrives as the Markdown the LLM produced; a small renderer maps
headings, bullets, **bold** and `code` onto Word styles.
"""

import io
import re

from docx import Document


def build_docx(kind: str, text: str, title: str) -> bytes:
    doc = Document()
    if kind == "transcript":
        _build_transcript(doc, text, title)
    else:
        _build_summary(doc, text, title)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _build_transcript(doc: Document, text: str, title: str) -> None:
    doc.add_heading(f"Transcript — {title or 'Lesson'}", 0)
    for line in text.splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        match = re.match(r"^\[([0-9:]+)\]\s*(.*)$", line)
        paragraph = doc.add_paragraph()
        if match:
            stamp = paragraph.add_run(f"[{match.group(1)}] ")
            stamp.bold = True
            _add_inline(paragraph, match.group(2))
        else:
            _add_inline(paragraph, line)


def _build_summary(doc: Document, text: str, title: str) -> None:
    doc.add_heading(f"Lesson Summary — {title or 'Lesson'}", 0)
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("### "):
            doc.add_heading(_plain(stripped[4:]), 2)
        elif stripped.startswith("## "):
            doc.add_heading(_plain(stripped[3:]), 1)
        elif stripped.startswith("# "):
            doc.add_heading(_plain(stripped[2:]), 1)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            paragraph = doc.add_paragraph(style="List Bullet")
            _add_inline(paragraph, stripped[2:])
        else:
            paragraph = doc.add_paragraph()
            _add_inline(paragraph, stripped)


def _add_inline(paragraph, markdown: str) -> None:
    """Renders **bold** and `code` spans into runs on the paragraph."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", markdown):
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
        elif part:
            paragraph.add_run(part)


def _plain(markdown: str) -> str:
    return re.sub(r"\*\*([^*]+)\*\*", r"\1", markdown).replace("`", "")
